"""Import a "PlugDrop Admissions Intelligence" research workbook (.xlsx) into the CRM.

The workbook is human research. It is imported the same way the research agent's output is:
as a completed research result per institution that a person reviews and approves before
anything reaches contacts or findings. Nothing is marked verified that the workbook does not
mark verified, and nothing is invented:

* emails are kept only when the workbook's Email Status says VERIFIED (published on a credible
  source); they are stored as "published", never as deliverability-verified;
* phones marked UNVERIFIED or DO NOT USE are dropped;
* conflicting or unnamed people are not turned into named contacts;
* statements labelled INFERENCE / REPORTED / UNVERIFIED stay inferences.

Re-importing is safe: an account whose rows have not changed since the last import is skipped.

CLI:  python -m app.research.workbook FILE.xlsx [--commit] [--user admin@example.com]
"""

from __future__ import annotations

import hashlib
import io
import json
import re
from dataclasses import dataclass, field
from datetime import UTC, datetime
from typing import Any

from sqlalchemy import select
from sqlalchemy.orm import Session

from app import audit
from app.models import (
    CLOSED_STAGES,
    Activity,
    Institution,
    Lead,
    ResearchContactCandidate,
    ResearchEvidence,
    ResearchJob,
    ResearchRun,
    User,
)
from app.normalize import domain_from_website, normalize_email, normalize_name
from app.research.reasoning import USE_CASES
from app.research.scoring import buying_role
from app.research.service import load_config, rescore_run
from app.services import contact_duplicates

PROVIDER = "workbook"
REQUIRED_SHEETS = ("Master_List", "Contact_Database", "Account_Intelligence", "Sales_Intelligence")
MAX_BYTES = 15 * 1024 * 1024
EMPTY = re.compile(r"^\s*(not found|unknown|none( found| identified)?|n/?a|not checked|-)?\s*(\(.*\))?\s*$", re.I)
EMAIL_RE = re.compile(r"[A-Za-z0-9._%+'-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
PHONE_RE = re.compile(r"\+?\d[\d\- ()]{6,}\d")
DOMAINISH = re.compile(r"(?:https?://)?((?:[a-z0-9-]+\.)+[a-z]{2,}(?:/[^\s;,()]*)?)", re.I)
QUALIFIED = re.compile(r"inference|reported|unverified|verify|\?|likely|appears|probabl|claims?", re.I)

TITLES: list[tuple[str, str | None]] = [
    (r"prof\.?\s*\(dr\.?\)", "Prof."),
    (r"prof\.?\s+dr\.?", "Prof."),
    (r"dr\.?\s*/\s*er\.?", "Dr."),
    (r"dr\.?\s*\(lt\.?\s*col\.?\)", "Dr."),
    (r"baba\s*\(dr\.?\)", "Baba"),
    (r"lt\.?\s*gen\.?", "Lt Gen"),
    (r"prof\.?", "Prof."),
    (r"dr\.?", "Dr."),
    (r"er\.?", "Er."),
    (r"mrs\.?", "Mrs."),
    (r"mr\.?", "Mr."),
    (r"ms\.?", "Ms."),
    (r"smt\.?", "Smt."),
    (r"shri|sh\.", "Shri"),
    (r"s\.", "Sardar"),
    (r"col\.?", "Col."),
    (r"brig\.?", "Brig."),
    (r"mahant", "Mahant"),
    (r"ca", None),  # chartered accountant: a qualification, not a form of address
]
USE_CASE_RULES = [
    (r"whatsapp", "whatsapp_engagement"),
    (r"missed[- ]call", "missed_call_recovery"),
    (r"voice|call|helpline|ivr", "voice_enquiry_automation"),
    (r"counsell", "counselling"),
    (r"routing|qualif|lead capture|pre-qualif", "enquiry_qualification"),
    (r"nurtur|follow[- ]?up", "follow_up"),
    (r"\bcrm\b|integrat", "crm_integration"),
]
SIGNAL_RULES = [
    (r"admissions? (are |is )?open|intake|application(s)? (open|close)|last date", "intake_period"),
    (
        r"(?:vacanc|hiring|recruiting|walk-in)[^.;]{0,60}counsell?or|counsell?ors?[^.;]{0,30}(?:vacanc|hiring|recruit)",
        "hiring_admissions",
    ),
    (r"new campus|campus (launch|opening)", "new_campus"),
    (r"new (course|programme|program)|launch(ed|es)? .*(course|programme|program)|co-branded program", "new_course"),
    # A CRM *change* (not merely having one): new, migrating, implementing, tender/RFP.
    (
        r"(?:new|migrat\w*|implement\w*|switch\w*|replac\w*|tender|rfp|evaluat\w*|select\w*)[^.;]{0,30}\bcrm\b|\bcrm\b[^.;]{0,20}(?:migration|implementation|tender|rfp|rollout|roll-out)",
        "crm_initiative",
    ),
    (r"digital transformation|erp (roll|implement)", "digital_transformation"),
    (r"campaign", "admissions_campaign"),
    (r"\d+\+?\s*(city )?(offices|campuses|centres|centers)|multi-campus|several campuses", "multi_campus"),
    (r"slow|delay|unanswered|no response|waiting", "response_delays"),
]


@dataclass
class Report:
    file: str
    research_date: str | None
    dry_run: bool
    accounts_in_file: int = 0
    institutions_created: int = 0
    institutions_updated: int = 0
    institutions_unchanged: int = 0
    leads_created: int = 0
    results_created: int = 0
    results_unchanged: int = 0
    evidence: int = 0
    people: int = 0
    emails_kept: int = 0
    emails_dropped: int = 0
    phones_dropped: int = 0
    people_not_imported: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    job_id: int | None = None

    def as_dict(self) -> dict[str, Any]:
        return {k: v for k, v in self.__dict__.items()}


# ---------------------------------------------------------------------------
# parsing helpers (pure functions; unit-tested)
# ---------------------------------------------------------------------------
def s(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def present(v: Any) -> bool:
    return bool(s(v)) and not EMPTY.match(s(v))


def to_url(text: str | None, fallback: str | None = None) -> str | None:
    """First web address in a source note ('aksuniversity.ac.in/contact; /Officers' -> https://aksuniversity.ac.in/contact)."""
    m = DOMAINISH.search(text or "")
    if m and "@" not in m.group(0):
        return "https://" + m.group(1).rstrip(".")
    return fallback


def parse_person(raw: Any) -> dict[str, Any]:
    """Split 'Prof. (Dr.) K.K. Dwivedi' into salutation / first / last name.

    Returns kind: person | role (a bracketed office such as '[Admissions Office]') | ambiguous
    (several names or a CONFLICT) | vacancy | empty."""
    text = s(raw)
    if not text or re.fullmatch(r"\[?\s*(not found|leadership - not found)\s*\]?", text, re.I):
        return {"kind": "empty"}
    if re.search(r"vacanc", text, re.I):
        return {"kind": "vacancy", "label": text.strip("[] ")}
    if text.startswith("["):
        label = re.split(r"\s+-\s+", text.strip("[] "))[0].strip()
        if re.search(r"not found|not identified|not captured|not named|no named", text, re.I) and not label:
            return {"kind": "empty"}
        return {"kind": "role", "label": label[:100], "raw": text.strip("[] ")}
    if text.upper().startswith("CONFLICT") or re.search(r"\s/\s|\svs\.?\s", text):
        names = re.sub(r"^conflict:\s*", "", text, flags=re.I)
        return {"kind": "ambiguous", "names": names}
    note = None
    m = re.search(r"\(([^)]*)\)\s*$", text)
    if m and not re.fullmatch(r"dr\.?|lt\.? col\.?|retd\.?", m.group(1), re.I):
        note = m.group(1)
        text = text[: m.start()].strip()
    text = re.sub(r",\s*(ias|ips|avsm.*|retd.*)$", "", text, flags=re.I)
    text = re.sub(r"\((retd|ct group)[^)]*\)", "", text, flags=re.I)
    text = re.sub(r"\s+ji$", "", text, flags=re.I).strip()
    salutation: str | None = None
    while True:
        for pat, label in TITLES:
            mm = re.match(rf"^(?:{pat})(?=\s)\s*", text, re.I)
            if mm:
                if salutation is None and label:
                    salutation = label
                text = text[mm.end() :]
                break
        else:
            break
    parts = text.split()
    if not parts:
        return {"kind": "empty"}
    first, last = (" ".join(parts[:-1]), parts[-1]) if len(parts) > 1 else (parts[0], "")
    return {"kind": "person", "salutation": salutation, "first_name": first[:100], "last_name": last[:100], "note": note}


def parse_email(value: Any, status: Any) -> tuple[str | None, str | None]:
    """(email, note). Kept only when the workbook marks it VERIFIED (published on a credible source)."""
    found = EMAIL_RE.findall(s(value))
    st = s(status)
    if not found:
        return None, None
    if not st.upper().startswith("VERIFIED"):
        return None, f"Email {found[0]} not imported: workbook status '{st or 'blank'}'."
    note = f"Workbook email status: {st}."
    if len(found) > 1:
        note += " Also listed: " + ", ".join(found[1:]) + "."
    return normalize_email(found[0]), note


def parse_phone(value: Any, phone_type: Any) -> tuple[str | None, str | None]:
    text, ptype = s(value), s(phone_type)
    if not present(text):
        return None, None
    if "DO NOT USE" in text.upper() or "UNVERIFIED" in ptype.upper() or "UNVERIFIED" in text.upper():
        return None, f"Phone not imported ({ptype or 'flagged'}): {text}"[:300]
    nums = [n.strip() for n in PHONE_RE.findall(text)]
    if not nums:
        return None, None
    note = f"Phone type: {ptype}." if ptype else None
    if len(nums) > 1 or len(text) > len(nums[0]) + 2:
        note = ((note + " ") if note else "") + f"All numbers: {text}"
    return nums[0][:50], note


def classify(text: str, rules: list[tuple[str, str]], default: str) -> str:
    for pat, cat in rules:
        if re.search(pat, text, re.I):
            return cat
    return default


def level_for(text: str) -> tuple[str, int]:
    if QUALIFIED.search(text):
        return "source_inference", 45
    if text.lower().startswith("confirmed"):
        return "verified_fact", 80
    return "source_inference", 65


# ---------------------------------------------------------------------------
# workbook reading
# ---------------------------------------------------------------------------
def read_workbook(data: bytes) -> dict[str, list[dict[str, Any]]]:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    missing = [n for n in REQUIRED_SHEETS if n not in wb.sheetnames]
    if missing:
        raise ValueError(f"This does not look like a research workbook: missing sheet(s) {', '.join(missing)}")
    out: dict[str, list[dict[str, Any]]] = {}
    for name in wb.sheetnames:
        rows = wb[name].iter_rows(values_only=True)
        if name == "README":
            out[name] = [{"text": s(r[0])} for r in rows if r and r[0]]
            continue
        header = [s(h) for h in next(rows, ())]
        out[name] = [dict(zip(header, r, strict=False)) for r in rows if r and any(c not in (None, "") for c in r)]
    wb.close()
    return out


def research_date(book: dict[str, list[dict[str, Any]]]) -> datetime | None:
    for r in book.get("README", []):
        m = re.search(r"research date:\s*(\d{1,2}-[A-Za-z]{3}-\d{4})", r["text"], re.I)
        if m:
            return datetime.strptime(m.group(1), "%d-%b-%Y").replace(tzinfo=UTC)
    return None


# ---------------------------------------------------------------------------
# import
# ---------------------------------------------------------------------------
def _city_from_address(address: str) -> str:
    part = [p.strip() for p in address.split(",") if p.strip()]
    last = re.sub(r"[\d\-]+|\bH\.?P\.?\b|\bU\.?P\.?\b|\bM\.?P\.?\b", "", part[-1] if part else "").strip(" .")
    return last[:100] if 2 < len(last) < 30 and re.fullmatch(r"[A-Za-z .()-]+", last) else ""


def _free_domain(db: Session, website: str | None, rep: Report, name: str) -> str | None:
    """Domains are unique in the CRM. A group domain shared by several universities is kept only on the first."""
    dom = domain_from_website(website)
    if dom and db.scalar(select(Institution.id).where(Institution.domain == dom)):
        rep.warnings.append(f"{name}: website domain {dom} is already used by another institution; kept the website only.")
        return None
    return dom


def _find_institution(db: Session, account: str, name: str) -> Institution | None:
    """Match on the workbook account id first (stable across re-imports), then on the exact normalised name.
    Websites are not used: several universities share a group domain (e.g. amity.edu)."""
    real = Institution.is_demo.is_(False)
    inst = db.scalar(select(Institution).where(real, Institution.notes.like(f"%Workbook account {account} %")))
    return inst or db.scalar(select(Institution).where(real, Institution.normalized_name == normalize_name(name)))


def _account_hash(account: str, inst_name: str, book: dict[str, list[dict[str, Any]]]) -> str:
    keyed = ("Account_Intelligence", "Contact_Database", "Data_Quality", "Sources")
    named = ("Sales_Intelligence", "Technology_Map", "Buying_Centre", "Contactability")
    rows: dict[str, Any] = {}
    for sh in keyed:
        rows[sh] = [r for r in book.get(sh, []) if s(r.get("Account ID")) == account]
    for sh in named:
        rows[sh] = [r for r in book.get(sh, []) if s(r.get("Institution")) == inst_name]
    return hashlib.sha256(json.dumps(rows, default=str, sort_keys=True).encode()).hexdigest()


def import_workbook(db: Session, user: User, data: bytes, filename: str, dry_run: bool = True) -> Report:
    if len(data) > MAX_BYTES:
        raise ValueError("File is larger than 15 MB")
    book = read_workbook(data)
    rdate = research_date(book)
    rep = Report(file=filename, research_date=rdate.date().isoformat() if rdate else None, dry_run=dry_run)
    when = rdate or datetime.now(UTC)
    if rdate is None:
        rep.warnings.append("No 'Research date' found in the README sheet; using today's date.")

    ai = {s(r.get("Account ID")): r for r in book["Account_Intelligence"] if s(r.get("Account ID"))}
    si = {s(r.get("Institution")): r for r in book["Sales_Intelligence"] if s(r.get("Institution"))}
    tech = {s(r.get("Institution")): r for r in book.get("Technology_Map", [])}
    sources: dict[str, list[str]] = {}
    for r in book.get("Sources", []):
        if s(r.get("URL")):
            sources.setdefault(s(r.get("Account ID")), []).append(s(r.get("URL")))
    quality: dict[str, list[str]] = {}
    for r in book.get("Data_Quality", []):
        if s(r.get("Issue")):
            quality.setdefault(s(r.get("Account ID")), []).append(s(r.get("Issue")))
    people: dict[str, list[dict[str, Any]]] = {}
    for r in book["Contact_Database"]:
        if s(r.get("Account ID")):
            people.setdefault(s(r.get("Account ID")), []).append(r)
    site_of = {a: to_url(s(r.get("Website"))) for a, r in ai.items() if present(r.get("Website"))}
    for a, rows in people.items():
        if a not in site_of and present(rows[0].get("Website")):
            site_of[a] = to_url(s(rows[0].get("Website")))

    job: ResearchJob | None = None
    cfg = load_config(db)

    master = [m for m in book["Master_List"] if s(m.get("Name of the University")) and s(m.get("Original Sr No")).isdigit()]
    name_count: dict[str, int] = {}
    for m in master:
        k = normalize_name(s(m.get("Name of the University")))
        name_count[k] = name_count.get(k, 0) + 1

    for m in master:
        # README: Account ID = "PD-" + original Sr No (pending rows often leave it blank)
        account = s(m.get("Account ID")) or f"PD-{int(s(m.get('Original Sr No'))):03d}"
        base = s(ai.get(account, {}).get("Institution")) or s(m.get("Name of the University"))
        name = base
        if normalize_name(name) == normalize_name(s(m.get("Name of the University"))) and (name_count[normalize_name(name)] > 1):
            name = f"{name} ({s(m.get('State')) or account})"  # same name, different university (one per state)
        rep.accounts_in_file += 1
        status = s(m.get("Research Status")) or "Pending"
        website = site_of.get(account) or None
        state, ugc, address, pin = s(m.get("State")), s(m.get("Status (UGC)")), s(m.get("Address")), s(m.get("Zip"))
        city = s(people[account][0].get("City")) if people.get(account) else _city_from_address(address)
        note_line = (
            f"Workbook account {account} ({status}). State: {state}. UGC status: {ugc}. Address: {address}"
            + (f", PIN {pin}" if pin else "")
            + "."
        )

        inst = _find_institution(db, account, name)
        if inst is None:
            inst = Institution(
                name=name[:300],
                normalized_name=normalize_name(name),
                website=website,
                domain=_free_domain(db, website, rep, name),
                institution_type="university",
                country="India",
                city=city,
                source=f"Research workbook: {filename}"[:200],
                notes=note_line,
            )
            db.add(inst)
            db.flush()
            audit.record(db, user, "create", "institution", inst.id, {**audit.snapshot(inst), "from_workbook": filename})
            rep.institutions_created += 1
        else:
            changes: dict[str, Any] = {}
            if not inst.website and website:
                inst.website, inst.domain = website, _free_domain(db, website, rep, name)
                changes["website"] = website
            if not inst.city and city:
                inst.city, changes["city"] = city, city
            if not inst.country:
                inst.country, changes["country"] = "India", "India"
            if account not in (inst.notes or ""):
                inst.notes = ((inst.notes + "\n") if inst.notes else "") + note_line
                changes["notes"] = "workbook account added"
            if changes:
                audit.record(db, user, "update", "institution", inst.id, {**changes, "from_workbook": filename})
                rep.institutions_updated += 1
            else:
                rep.institutions_unchanged += 1

        if not status.lower().startswith("researched") or account not in ai and account not in people:
            continue

        on_hold = "hold" in status.lower()
        acct_hash = _account_hash(account, base, book)
        same = db.scalar(
            select(ResearchRun.id).where(
                ResearchRun.institution_id == inst.id,
                ResearchRun.summary["workbook"]["hash"].astext == acct_hash,
            )
        )
        if same:
            rep.results_unchanged += 1
            continue
        busy = db.scalar(
            select(ResearchRun.id).where(ResearchRun.institution_id == inst.id, ResearchRun.status.in_(["queued", "running"]))
        )
        if busy:
            rep.warnings.append(f"{name}: research is queued or running in the app; workbook data skipped for now.")
            continue

        if job is None:
            job = ResearchJob(
                label=f"Workbook import: {filename}"[:200],
                provider=PROVIDER,
                reasoner="human",
                requested_by_id=user.id,
                status="completed",
                started_at=datetime.now(UTC),
                finished_at=datetime.now(UTC),
            )
            db.add(job)
            db.flush()
            rep.job_id = job.id

        acc, sales, tm = ai.get(account, {}), si.get(base, {}), tech.get(base, {})
        main_src = to_url(sources.get(account, [None])[0]) or to_url(s(acc.get("Source"))) or website
        run = ResearchRun(
            job_id=job.id,
            institution_id=inst.id,
            status="succeeded",
            attempts=1,
            started_at=when,
            finished_at=when,
            log=[{"at": datetime.now(UTC).isoformat(), "event": "imported", "file": filename, "account": account}],
        )
        db.add(run)
        db.flush()
        job.total_items += 1

        def add_ev(
            kind: str,
            category: str,
            claim: str,
            src: str | None,
            lvl: tuple[str, int] | None = None,
            *,
            run: ResearchRun = run,
            inst: Institution = inst,  # type: ignore[assignment]
            account: str = account,
        ) -> None:
            level, conf = lvl or level_for(claim)
            if not src:
                level = "hypothesis"
            db.add(
                ResearchEvidence(
                    run_id=run.id,
                    institution_id=inst.id,
                    kind=kind,
                    category=category,
                    claim=claim[:2000],
                    evidence_level=level,
                    confidence=conf,
                    source_url=src,
                    source_title=f"Research workbook {account}" if src else None,
                    retrieved_at=when,
                    provider=PROVIDER,
                )
            )
            rep.evidence += 1

        if website:
            add_ev("profile", "website", f"Official website: {domain_from_website(website)}", website, ("verified_fact", 85))
        if city or state:
            add_ev("profile", "location", f"Located in {', '.join(x for x in (city, state, 'India') if x)}", website or main_src)
        triggers = [t.strip() for t in re.split(r";", s(acc.get("Sales Trigger")) or s(sales.get("Trigger"))) if present(t)]
        for t in triggers:
            add_ev("signal", classify(t, SIGNAL_RULES, "other"), f"Trigger: {t}", main_src)
        if present(sales.get("Current Situation")):
            t = s(sales.get("Current Situation"))
            add_ev("signal", classify(t, SIGNAL_RULES[-2:], "other"), t, main_src)
        if present(sales.get("Potential Gap")):
            t = s(sales.get("Potential Gap"))
            add_ev("signal", classify(t, SIGNAL_RULES[-1:], "other"), f"Possible gap: {t}", main_src)
        tech_seen = [
            f"{k}: {s(v)}"
            for k, v in (tm or acc).items()
            if k in ("CRM", "ERP/SIS", "WhatsApp", "Chatbot", "Voice/IVR", "AI", "Marketing Automation", "Call Centre")
            and re.match(r"(confirmed|evidence found)", s(v), re.I)
        ]
        if tech_seen:
            add_ev("signal", "other", "Technology seen: " + "; ".join(tech_seen), main_src)
        for p in people.get(account, []):
            pp = parse_person(p.get("Person"))
            if pp["kind"] == "vacancy":
                cat = "hiring_admissions" if re.search(r"counsell?or", pp["label"], re.I) else "other"
                add_ev("signal", cat, f"Vacancy: {pp['label']}", to_url(s(p.get("Person Source")), main_src))
        relevance = s(sales.get("PlugDrop Relevance"))
        primary = classify(relevance, USE_CASE_RULES, "voice_enquiry_automation") if present(relevance) else None
        if primary:
            add_ev("use_case", primary, f"PlugDrop relevance (researcher): {relevance}", main_src, ("source_inference", 60))

        missing = [
            {"item": q, "how_to_verify": "Flagged in the workbook's Data_Quality sheet."} for q in quality.get(account, [])
        ]
        for p in people.get(account, []):
            rep_person = _add_candidate(db, run, inst, p, account, main_src or website, when, rep)
            if rep_person:
                missing.append(rep_person)
        db.flush()
        db.refresh(run)
        run.summary = {
            "fit_summary": s(sales.get("Current Situation")) or "See imported evidence.",
            "primary_use_case": primary,
            "primary_use_case_label": USE_CASES.get(primary or ""),
            "use_case_rationale": f"Researcher's view of PlugDrop relevance: {relevance}" if relevance else "",
            "missing_information": missing,
            "profile_confirmed": bool(website),
            "campus_count": None,
            "documents_reviewed": len(sources.get(account, [])),
            "best_entry_point": s(sales.get("Best Entry Point")) or None,
            "opportunity_signal": s(acc.get("Opportunity Signal") or sales.get("Opportunity Signal")) or None,
            "digital_maturity": s(acc.get("Digital Maturity")) or None,
            "admissions_model": s(sales.get("Admissions Model")) or None,
            "on_hold": on_hold,
            "sources": sources.get(account, [])[:20],
            "workbook": {"file": filename, "account_id": account, "hash": acct_hash, "research_date": rep.research_date},
        }
        rescore_run(db, run, cfg)
        signal = run.summary.get("opportunity_signal")
        if signal or run.summary.get("best_entry_point"):
            # The score counts only evidence the CRM can classify; show the researcher's own judgement next to it.
            run.summary = {
                **run.summary,
                "next_action": (
                    f"Researcher's rating: {signal or 'not given'}. "
                    f"Suggested entry point: {run.summary.get('best_entry_point') or 'not given'}. "
                    f"(CRM score suggestion: {run.summary.get('next_action')})"
                ),
            }
        inst.research_score, inst.research_tier, inst.last_researched_at = run.score, run.tier, when
        rep.results_created += 1

        lead = db.scalar(select(Lead).where(Lead.institution_id == inst.id, Lead.stage.not_in([c.value for c in CLOSED_STAGES])))
        if on_hold:
            rep.warnings.append(f"{name} is marked ON HOLD in the workbook: research imported, no lead created.")
        elif lead is None:
            lead = Lead(
                institution_id=inst.id,
                stage="researched",
                score=run.score or 0,
                source="Research workbook",
                summary=(
                    f"Opportunity signal: {run.summary['opportunity_signal'] or 'n/a'}. "
                    f"Best entry point: {run.summary['best_entry_point'] or 'n/a'}"
                )[:2000],
            )
            db.add(lead)
            db.flush()
            audit.record(db, user, "create", "lead", lead.id, {"institution_id": inst.id, "from_workbook": filename})
            rep.leads_created += 1
        db.add(
            Activity(
                activity_type="research",
                subject=f"Research imported from workbook ({account}): score {run.score}, tier {run.tier}",
                body=f"{rep.research_date or ''} research by the team. Review and approve it on the research result page.",
                institution_id=inst.id,
                lead_id=lead.id if lead else None,
                user_id=user.id,
                details={"research_run_id": run.id, "workbook": filename},
            )
        )

    if job is not None:
        audit.record(db, user, "import", "research_job", job.id, {k: v for k, v in rep.as_dict().items() if k != "warnings"})
    if dry_run:
        db.rollback()
        rep.job_id = None
    else:
        db.commit()
    return rep


def _add_candidate(
    db: Session,
    run: ResearchRun,
    inst: Institution,
    p: dict[str, Any],
    account: str,
    fallback_src: str | None,
    when: datetime,
    rep: Report,
) -> dict[str, str] | None:
    """Create one reviewable contact candidate. Returns a 'missing information' item when the row can't be a contact."""
    pp = parse_person(p.get("Person"))
    designation = s(p.get("Designation"))[:200] or None
    email, email_note = parse_email(p.get("Professional Email"), p.get("Email Status"))
    phone, phone_note = parse_phone(p.get("Direct Business Phone"), p.get("Phone Type"))
    if email_note and not email:
        rep.emails_dropped += 1
    if phone_note and not phone:
        rep.phones_dropped += 1
    label = f"{inst.name}: {s(p.get('Person')) or designation}"
    if pp["kind"] in ("empty", "vacancy"):
        return None
    if pp["kind"] == "ambiguous":
        if not (email or phone):
            rep.people_not_imported.append(f"{label} (conflicting names, no contact detail)")
            return {
                "item": f"Who is the current {designation or 'office holder'}? Sources disagree: {pp['names']}",
                "how_to_verify": "Check the institution's official authorities / leadership page.",
            }
        pp = {"kind": "role", "label": (designation or "Office")[:100], "raw": pp["names"]}
    if pp["kind"] == "role" and not (email or phone):
        rep.people_not_imported.append(f"{label} (no named person and no contact detail)")
        return None
    src = to_url(s(p.get("Person Source")), None) or to_url(s(p.get("Email Source")), None) or fallback_src
    if not src:
        rep.people_not_imported.append(f"{label} (no source)")
        return None
    notes = [
        f"Workbook {account}",
        f"confidence {s(p.get('Confidence')) or '?'}",
        f"role verified: {s(p.get('Current Role Verified')) or '?'}",
    ]
    if present(p.get("Buying Role")):
        notes.append(f"buying role: {s(p.get('Buying Role'))}")
    if present(p.get("Contact Priority")):
        notes.append(f"priority {s(p.get('Contact Priority'))}")
    if pp["kind"] == "role":
        notes.append(f"office/role mailbox: {pp.get('raw')}")
    if pp.get("note"):
        notes.append(pp["note"])
    if present(p.get("Person Source")):
        notes.append(f"source: {s(p.get('Person Source'))}")
    text = "; ".join(notes) + "."
    for extra in (email_note, phone_note, s(p.get("Notes")) or None):
        if extra:
            text += " " + extra
    is_person = pp["kind"] == "person"
    first = pp["first_name"] if is_person else pp["label"]
    last = pp["last_name"] if is_person else ""
    role, _ = buying_role(designation)
    dups = [m for m in contact_duplicates(db, inst.id, first, last, email, None) if m.exact]
    db.add(
        ResearchContactCandidate(
            run_id=run.id,
            institution_id=inst.id,
            salutation=pp.get("salutation") if is_person else None,
            first_name=first,
            last_name=last,
            designation=designation,
            department=s(p.get("Department"))[:200] or None,
            buying_role=role,
            email=email,
            email_status="published" if email else "none",
            phone=phone,
            verification_status="unverified",
            source_url=src[:1000],
            retrieved_at=when,
            provider=PROVIDER,
            notes=text[:500],
            existing_contact_id=dups[0].id if dups else None,
        )
    )
    rep.people += 1
    if email:
        rep.emails_kept += 1
    return None


def main() -> None:  # pragma: no cover - thin CLI
    import argparse
    import pathlib

    from app.db import SessionLocal

    ap = argparse.ArgumentParser(description="Import a PlugDrop research workbook (.xlsx)")
    ap.add_argument("file")
    ap.add_argument("--commit", action="store_true", help="Save to the database (default: preview only)")
    ap.add_argument("--user", default=None, help="Email of the admin/manager recorded as the importer")
    args = ap.parse_args()
    path = pathlib.Path(args.file)
    with SessionLocal() as db:
        q = select(User).where(User.role.in_(["admin", "manager"]), User.is_active.is_(True))
        if args.user:
            q = q.where(User.email == args.user.lower())
        user = db.scalars(q.order_by(User.id)).first()
        if user is None:
            raise SystemExit("No matching admin/manager user found")
        rep = import_workbook(db, user, path.read_bytes(), path.name, dry_run=not args.commit)
    print(json.dumps(rep.as_dict(), indent=2))


if __name__ == "__main__":  # pragma: no cover
    main()
