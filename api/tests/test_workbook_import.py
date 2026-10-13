"""Import of the team's research workbook (.xlsx) into research results that people review before they reach the CRM."""

import io

import openpyxl
import pytest
from sqlalchemy import func, select

from app.db import SessionLocal
from app.models import Contact, Institution, Lead, ResearchContactCandidate, ResearchEvidence, ResearchRun
from app.research.workbook import parse_email, parse_person, parse_phone, to_url

# ---------------------------------------------------------------------------
# parsing rules
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "raw,sal,first,last",
    [
        ("Prof. B.A. Chopade", "Prof.", "B.A.", "Chopade"),
        ("Prof. (Dr.) K.K. Dwivedi", "Prof.", "K.K.", "Dwivedi"),
        ("Mr. Sonu Kumar Soni", "Mr.", "Sonu Kumar", "Soni"),
        ("Dr. V.K. Sharma (Virendra Kumar Sharma)", "Dr.", "V.K.", "Sharma"),
        ("Lt Gen V K Sharma, AVSM (Retd)", "Lt Gen", "V K", "Sharma"),
        ("Mr. Amit Aggarwal, IAS", "Mr.", "Amit", "Aggarwal"),
        ("CA Rajat Bhatia", None, "Rajat", "Bhatia"),
        ("Alka Das", None, "Alka", "Das"),
        ("Dr./Er. L.K. Abhilashi", "Dr.", "L.K.", "Abhilashi"),
        ("S. Rashpal Singh Dhaliwal", "Sardar", "Rashpal Singh", "Dhaliwal"),
    ],
)
def test_parse_person_names(raw, sal, first, last):
    p = parse_person(raw)
    assert p["kind"] == "person"
    assert (p["salutation"], p["first_name"], p["last_name"]) == (sal, first, last)


def test_parse_person_special_rows():
    assert parse_person("[Admissions Office - no named head]") == {
        "kind": "role",
        "label": "Admissions Office",
        "raw": "Admissions Office - no named head",
    }
    assert parse_person("CONFLICT: Dr. R.K. Pandey / Prof. Kalyan K. Sahoo")["kind"] == "ambiguous"
    assert parse_person("Mr. Karambir Singh / Mr. Kamal Pama")["kind"] == "ambiguous"
    assert parse_person("[Director - Admissions and Marketing - VACANCY]")["kind"] == "vacancy"
    assert parse_person("[Not Found]")["kind"] == "empty"


def test_only_verified_emails_are_kept_and_never_marked_deliverable():
    assert parse_email("admission@aksuniversity.com", "VERIFIED (role mailbox)")[0] == "admission@aksuniversity.com"
    email, note = parse_email("registrar@adeshuniversity.ac.in (role mailbox)", "VERIFIED (role mailbox, AIU directory)")
    assert email == "registrar@adeshuniversity.ac.in" and "AIU directory" in note
    email, note = parse_email("admissions@x.edu.in; info@x.edu.in", "VERIFIED (official contact page)")
    assert email == "admissions@x.edu.in" and "info@x.edu.in" in note
    # unverified / not relied on -> dropped, with a reason
    assert parse_email("info@abhilashiuniversity.ac.in", "UNVERIFIED (3rd-party listing)")[0] is None
    assert parse_email("Not Found", "Not Found") == (None, None)
    assert parse_email("vc@x.in", "Not Found (vc@ address in PDF appears mistyped - not used)")[0] is None


def test_phone_rules():
    assert parse_phone("09049483914", "Official Campus") == ("09049483914", "Phone type: Official Campus.")
    phone, note = parse_phone("75270-11305 (mob); 0164-5055204", "Official Department")
    assert phone == "75270-11305" and "0164-5055204" in note
    assert parse_phone("DO NOT USE 9818333993 (listed for predecessor)", "DO NOT USE")[0] is None
    assert parse_phone("9999999999", "UNVERIFIED (third-party listing)")[0] is None
    assert parse_phone("Not Found", None) == (None, None)


def test_to_url():
    assert to_url("aksuniversity.ac.in/contact; /Officers-of-University") == "https://aksuniversity.ac.in/contact"
    assert to_url("en.wikipedia.org (REPORTED)") == "https://en.wikipedia.org"
    assert to_url("Public Self-Disclosure Jul-2026", "https://x.ac.in") == "https://x.ac.in"


# ---------------------------------------------------------------------------
# end-to-end import through the API
# ---------------------------------------------------------------------------
CONTACT_COLS = [
    "Account ID", "Original Sr No", "Institution", "University Group", "Campus", "City", "State", "Website", "Person",
    "Designation", "Department", "Buying Role", "Contact Priority", "Professional Email", "Email Status",
    "Direct Business Phone", "Phone Type", "WhatsApp", "LinkedIn", "Current Role Verified", "Person Source",
    "Email Source", "Phone Source", "LinkedIn Source", "Confidence", "Notes",
]  # fmt: skip


def workbook_bytes(trigger: str = "Admissions Open 2026-27") -> bytes:
    wb = openpyxl.Workbook()
    wb.active.title = "README"
    wb["README"].append(["PlugDrop.AI - Indian Higher-Ed Admissions Intelligence Database"])
    wb["README"].append(["Researched so far: rows 1-2. Research date: 30-Sep-2026."])
    ml = wb.create_sheet("Master_List")
    ml.append(["Original Sr No", "Type", "Name of the University", "Address", "Zip", "State", "Status (UGC)", "Account ID", "Research Status"])  # fmt: skip
    ml.append([1, "Private", "Testpur University", "Main Road, Testpur", 111111, "Gujarat", "2(f)", "PD-001", "Researched"])  # fmt: skip
    ml.append([2, "Private", "Sample Vidyapeeth", "Sample Nagar", 222222, "Punjab", "2(f)", None, "Pending"])
    ml.append([3, "Private", "Twin University", "Jaipur", 333333, "Rajasthan", "2(f)", None, "Pending"])
    ml.append([4, "Private", "Twin University", "Gwalior", 444444, "Madhya Pradesh", "2(f)", None, "Pending"])
    cd = wb.create_sheet("Contact_Database")
    cd.append(CONTACT_COLS)
    base = ["PD-001", 1, "Testpur University", None, "Main", "Testpur", "Gujarat", "https://testpur.ac.in"]
    tail = ["Not Found", "Not Found (not checked)", "Yes", "testpur.ac.in/authorities", "same", "same", "Not checked"]
    cd.append(base + ["Prof. Asha Rao", "Vice-Chancellor", "Leadership", "Executive Sponsor", 2, "vc@testpur.ac.in", "VERIFIED (role mailbox)", "0261-1234567", "Official Campus"] + tail + ["High", None])  # fmt: skip
    cd.append(base + ["Mr. Ravi Kumar", "Director - Admissions", "Admissions", "Business Owner", 1, "ravi@testpur.ac.in", "UNVERIFIED (3rd-party listing)", "DO NOT USE 99999", "DO NOT USE"] + tail + ["Low", None])  # fmt: skip
    cd.append(base + ["[Admissions Office]", "Admissions Office", "Admissions", "Business Owner", 1, "admissions@testpur.ac.in", "VERIFIED (official site)", "Not Found", None] + tail + ["Medium", None])  # fmt: skip
    cd.append(base + ["CONFLICT: Dr. A / Dr. B", "Registrar", "Administration", "Procurement", 3, "Not Found", "Not Found", "Not Found", None] + tail + ["Low", None])  # fmt: skip
    ai = wb.create_sheet("Account_Intelligence")
    ai.append(["Account ID", "Institution", "Campus", "Website", "CRM", "ERP/SIS", "WhatsApp", "Call Centre", "Digital Maturity", "Sales Trigger", "Opportunity Signal", "Source"])  # fmt: skip
    ai.append(["PD-001", "Testpur University", "Main", "https://testpur.ac.in", "Unknown", "Confirmed: in-house ERP", "Not Found", "Evidence Found (helplines)", "Moderate", trigger, "STRONG", "testpur.ac.in/contact"])  # fmt: skip
    si = wb.create_sheet("Sales_Intelligence")
    si.append(["Institution", "Current Situation", "Admissions Model", "Technology", "Potential Gap", "PlugDrop Relevance", "Trigger", "Best Entry Point", "Opportunity Signal"])  # fmt: skip
    si.append(["Testpur University", "Admissions run from 12 city offices plus a helpline.", "Walk-in + online", "ERP", "INFERENCE: enquiries from offices likely not unified.", "WhatsApp follow-up and lead routing", trigger, "Registrar", "STRONG"])  # fmt: skip
    dq = wb.create_sheet("Data_Quality")
    dq.append(["Account ID", "Issue"])
    dq.append(["PD-001", "Two email domains in use"])
    src = wb.create_sheet("Sources")
    src.append(["Account ID", "URL"])
    src.append(["PD-001", "https://testpur.ac.in/contact"])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def upload(client, data: bytes, dry_run: bool):
    return client.post(
        f"/api/research/import-workbook?dry_run={'true' if dry_run else 'false'}",
        files={"file": ("intel.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )


def test_reps_cannot_import(rep):
    assert upload(rep, workbook_bytes(), True).status_code == 403


def test_rejects_non_workbooks(manager):
    r = manager.post("/api/research/import-workbook", files={"file": ("x.xlsx", b"not a zip", "application/octet-stream")})
    assert r.status_code == 422
    r = manager.post("/api/research/import-workbook", files={"file": ("x.csv", b"a,b", "text/csv")})
    assert r.status_code == 422


def test_dry_run_saves_nothing(manager):
    r = upload(manager, workbook_bytes(), True)
    assert r.status_code == 200, r.text
    rep = r.json()
    assert rep["dry_run"] is True and rep["accounts_in_file"] == 4 and rep["results_created"] == 1
    with SessionLocal() as db:
        assert db.scalar(select(func.count()).select_from(Institution).where(Institution.is_demo.is_(False))) == 0


def test_import_creates_reviewable_research_and_respects_email_rules(manager):
    r = upload(manager, workbook_bytes(), False)
    assert r.status_code == 200, r.text
    rep = r.json()
    assert rep["institutions_created"] == 4 and rep["leads_created"] == 1 and rep["results_created"] == 1
    assert rep["emails_kept"] == 2 and rep["emails_dropped"] == 1 and rep["phones_dropped"] == 1
    assert any("CONFLICT" in p for p in rep["people_not_imported"])
    with SessionLocal() as db:
        names = set(db.scalars(select(Institution.name).where(Institution.is_demo.is_(False))))
        # same name in two states -> two institutions, told apart by state; pending rows get PD-<Sr No>
        assert {"Twin University (Rajasthan)", "Twin University (Madhya Pradesh)"} <= names
        inst = db.scalar(select(Institution).where(Institution.name == "Testpur University"))
        assert inst.country == "India" and inst.domain == "testpur.ac.in" and "PD-001" in inst.notes
        assert db.scalar(select(Lead.stage).where(Lead.institution_id == inst.id)) == "researched"
        run = db.scalar(select(ResearchRun).where(ResearchRun.institution_id == inst.id))
        assert run.status == "succeeded" and run.score is not None and run.summary["primary_use_case"] == "whatsapp_engagement"
        assert run.summary["opportunity_signal"] == "STRONG"
        assert any("Two email domains" in m["item"] for m in run.summary["missing_information"])
        assert any("Registrar" in m["item"] for m in run.summary["missing_information"])
        ev = {e.category: e for e in db.scalars(select(ResearchEvidence).where(ResearchEvidence.run_id == run.id))}
        assert ev["intake_period"].evidence_level == "source_inference" and ev["intake_period"].source_url
        assert ev["multi_campus"].claim.startswith("Admissions run from 12 city offices")
        assert all(e.review_status == "pending" for e in ev.values())
        cands = {
            c.first_name: c for c in db.scalars(select(ResearchContactCandidate).where(ResearchContactCandidate.run_id == run.id))
        }
        assert cands["Asha"].salutation == "Prof." and cands["Asha"].email == "vc@testpur.ac.in"
        assert cands["Asha"].email_status == "published" and cands["Asha"].verification_status == "unverified"
        assert cands["Ravi"].email is None and cands["Ravi"].phone is None and "not imported" in cands["Ravi"].notes
        assert cands["Admissions Office"].email == "admissions@testpur.ac.in" and cands["Admissions Office"].last_name == ""
        assert db.scalar(select(func.count()).select_from(Contact)) == 0  # nothing reaches contacts before review
        run_id = run.id

    # A person approves and pushes; the contact keeps its salutation and notes
    detail = manager.get(f"/api/research/runs/{run_id}").json()
    asha = next(c for c in detail["candidates"] if c["first_name"] == "Asha")
    assert (
        manager.post(f"/api/research/runs/{run_id}/review", json={"candidates": {str(asha["id"]): "approved"}}).status_code == 200
    )
    assert manager.post(f"/api/research/runs/{run_id}/push").json()["contacts_created"] == 1
    with SessionLocal() as db:
        c = db.scalar(select(Contact))
        assert c.greeting_name == "Prof. Rao" and "confidence High" in c.notes

    # Re-importing the same workbook changes nothing
    again = upload(manager, workbook_bytes(), False).json()
    assert again["results_created"] == 0 and again["results_unchanged"] == 1 and again["institutions_created"] == 0
    # A newer workbook with changed research for the account creates a new result to review
    newer = upload(manager, workbook_bytes(trigger="New B.Tech programme launched for 2027"), False).json()
    assert newer["results_created"] == 1 and newer["leads_created"] == 0
